from pathlib import Path
import json
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.local_pipeline.io import read_df, resolve_table
from src.output.contracts import CALL_LIST_COLUMNS

qa_path = ROOT/"data"/"qa"/"business_output_qa.json"
qa = json.loads(qa_path.read_text(encoding="utf-8"))

main_path = Path(qa["paths"]["main_workbook"])
audit_path = Path(qa["paths"]["audit_workbook"])
cdata_path = Path(qa["paths"]["cdata_snapshot"])

main_wb = load_workbook(main_path, read_only=False, data_only=False)
audit_wb = load_workbook(audit_path, read_only=False, data_only=False)

ws = main_wb["Call List"]
headers = [ws.cell(1, c).value for c in range(1, len(CALL_LIST_COLUMNS) + 1)]
excel_rows = max(0, ws.max_row - 1)

risk_scores = [
    ws.cell(r, 8).value
    for r in range(2, ws.max_row + 1)
    if ws.cell(r, 8).value is not None
]
risk_sorted = risk_scores == sorted(risk_scores, reverse=True)

scores = read_df(
    resolve_table(
        ROOT/"data"/"scoring"/"dex_v2_category_loss_scores_wide_zipcode"
    )
)
scores["snapshot_dt"] = pd.to_datetime(scores["snapshot_dt"])
latest = scores["snapshot_dt"].max()
expected_rows = int((scores["snapshot_dt"] == latest).sum())

required_audit_sheets = {
    "Owner Mapping Review",
    "Customer ID Validation",
    "Validation Layer 1-3",
    "Row Flags",
    "Dormant Accounts",
    "Pipeline Health",
    "Reject Exclusions",
    "Customer Attributes",
}

checks = {
    "qa_status_pass": qa["status"] == "PASS",
    "main_workbook_exists": main_path.exists(),
    "audit_workbook_exists": audit_path.exists(),
    "cdata_snapshot_exists": cdata_path.exists(),
    "main_has_call_list": "Call List" in main_wb.sheetnames,
    "main_has_column_descriptions": "Column Descriptions" in main_wb.sheetnames,
    "main_has_hidden_lists": (
        "Lists" in main_wb.sheetnames
        and main_wb["Lists"].sheet_state == "hidden"
    ),
    "call_list_headers_exact": headers == CALL_LIST_COLUMNS,
    "call_list_rows_match_qa": excel_rows == int(qa["rows"]["call_list"]),
    "call_list_rows_match_scores": excel_rows == expected_rows,
    "risk_sorted_descending": risk_sorted,
    "all_audit_sheets_present": required_audit_sheets <= set(audit_wb.sheetnames),
    "all_output_validations_pass": all(
        value == "PASS" for value in qa["validation"].values()
    ),
}

result = {
    "status": "PASS" if all(checks.values()) else "FAIL",
    "checks": checks,
    "snapshot_dt": qa["snapshot_dt"],
    "rows": qa["rows"],
    "totals": qa["totals"],
    "paths": qa["paths"],
}

payload = json.dumps(result, indent=2)
(ROOT/"data"/"qa"/"business_output_validation.json").write_text(
    payload,
    encoding="utf-8",
)
print(payload)

if result["status"] != "PASS":
    raise SystemExit(1)
