from __future__ import annotations

from pathlib import Path
import json

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.features.contracts import CATEGORY_KEYS, CUSTOMER_COLUMN
from src.local_pipeline.io import read_df, resolve_table, write_df, update_manifest
from .contracts import CALL_LIST_COLUMNS, COLUMN_DESCRIPTIONS, VALID_ACTIONS


HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
GREY_FONT = "808080"
RED_FILL = "F4CCCC"
YELLOW_FILL = "FFF2CC"
GREEN_FILL = "D9EAD3"


def _read_optional(base: Path) -> pd.DataFrame:
    try:
        return read_df(resolve_table(base))
    except FileNotFoundError:
        return pd.DataFrame()


def _peak_text(row: pd.Series) -> str:
    vals = []
    for c in ["recent_peak_1", "recent_peak_2", "recent_peak_3"]:
        v = row.get(c)
        if pd.notna(v) and str(v).strip():
            vals.append(str(v))
    return ", ".join(vals) if vals else "No recent peak-month history"


def _profile_text(row: pd.Series) -> str:
    group = str(row.get("customergroupname_current") or "Unknown group")
    tenure_days = pd.to_numeric(
        pd.Series([row.get("tenure_days_all")]),
        errors="coerce",
    ).iloc[0]
    tenure = (
        f"{float(tenure_days) / 365.25:.1f} years"
        if pd.notna(tenure_days)
        else "tenure unavailable"
    )
    credit = pd.to_numeric(
        pd.Series([row.get("creditmax_current")]),
        errors="coerce",
    ).iloc[0]
    credit_text = ("$" + f"{float(credit):,.0f}") if pd.notna(credit) else "unavailable"
    peaks = _peak_text(row)
    top = row.get("top_3_categories")
    top_text = str(top) if pd.notna(top) and str(top).strip() else "unavailable"
    return (
        f"{group} customer for {tenure}. "
        f"Credit limit: {credit_text}. "
        f"Peak buying months: {peaks}. "
        f"Top categories: {top_text}."
    )


def _signal_summary(reason_long: pd.DataFrame) -> pd.DataFrame:
    if reason_long.empty:
        return pd.DataFrame(
            columns=CATEGORY_KEYS + ["bold_signals", "grey_signals"]
        )

    x = reason_long.copy()
    x["abs_shap_value"] = pd.to_numeric(
        x["abs_shap_value"],
        errors="coerce",
    ).fillna(0.0)
    x = x.sort_values(
        CATEGORY_KEYS + ["abs_shap_value", "reason_rank"],
        ascending=[True] * len(CATEGORY_KEYS) + [False, True],
    )

    rows = []
    for key, g in x.groupby(CATEGORY_KEYS, dropna=False, sort=False):
        kt = key if isinstance(key, tuple) else (key,)
        risk_up = g[g["direction"].astype(str).eq("risk_up")].head(2)
        bold = risk_up["reason_text"].dropna().astype(str).tolist()
        if not bold:
            bold = ["No strong risk-up SHAP driver in the leading features."]

        used_idx = set(risk_up.index)
        grey = (
            g[~g.index.isin(used_idx)]
            .head(3)["reason_text"]
            .dropna()
            .astype(str)
            .tolist()
        )
        if not grey:
            grey = ["No additional SHAP context available."]

        row = {c: v for c, v in zip(CATEGORY_KEYS, kt)}
        row["bold_signals"] = "\n".join(bold)
        row["grey_signals"] = "\n".join(grey)
        rows.append(row)
    return pd.DataFrame(rows)


def _make_call_list(
    scores: pd.DataFrame,
    reasons: pd.DataFrame,
    attrs: pd.DataFrame,
    peak: pd.DataFrame,
    top: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    scores = scores.copy()
    scores["snapshot_dt"] = pd.to_datetime(scores["snapshot_dt"])
    latest = scores["snapshot_dt"].max()
    scores = scores[scores["snapshot_dt"] == latest].copy()

    attr_cols = [
        CUSTOMER_COLUMN,
        "zipcode",
        "customer_ids",
        "customergroupname",
        "creditmax",
        "salesrep",
        "servicingbranch",
    ]
    have_attr = [c for c in attr_cols if c in attrs.columns]
    attrs2 = (
        attrs[have_attr].drop_duplicates([CUSTOMER_COLUMN, "zipcode"])
        if not attrs.empty
        else pd.DataFrame()
    )

    if not attrs2.empty:
        scores = scores.merge(
            attrs2,
            on=[CUSTOMER_COLUMN, "zipcode"],
            how="left",
            validate="many_to_one",
            suffixes=("", "_attr"),
        )

    if not peak.empty:
        scores = scores.merge(
            peak,
            on=[CUSTOMER_COLUMN, "zipcode"],
            how="left",
            validate="many_to_one",
        )
    if not top.empty:
        scores = scores.merge(
            top,
            on=[CUSTOMER_COLUMN, "zipcode"],
            how="left",
            validate="many_to_one",
        )

    sig = _signal_summary(reasons)
    if not sig.empty:
        scores = scores.merge(
            sig,
            on=CATEGORY_KEYS,
            how="left",
            validate="one_to_one",
        )

    scores["customer_id_output"] = scores.get(
        "customer_ids",
        pd.Series(index=scores.index, dtype="string"),
    )
    scores["customer_group_output"] = scores.get(
        "customergroupname_current",
        scores.get(
            "customergroupname",
            pd.Series(index=scores.index, dtype="string"),
        ),
    )
    scores["branch_output"] = scores.get(
        "servicingbranch",
        pd.Series(index=scores.index, dtype="string"),
    )
    scores["sales_rep_output"] = scores.get(
        "salesrep_current",
        scores.get(
            "salesrep",
            pd.Series(index=scores.index, dtype="string"),
        ),
    )

    scores["category_risk_score"] = np.rint(
        pd.to_numeric(scores["p_14"], errors="coerce") * 100
    ).astype("Int64")
    scores["category_annual_spend"] = pd.to_numeric(
        scores["spend_365"],
        errors="coerce",
    ).fillna(0.0)
    scores["category_revenue_at_risk"] = (
        scores["category_annual_spend"]
        * pd.to_numeric(scores["p_14"], errors="coerce").fillna(0.0)
    )
    scores["who_they_are"] = scores.apply(_profile_text, axis=1)
    scores["bold_signals"] = scores.get(
        "bold_signals",
        pd.Series(index=scores.index, dtype="string"),
    ).fillna("No strong risk-up SHAP driver in the leading features.")
    scores["grey_signals"] = scores.get(
        "grey_signals",
        pd.Series(index=scores.index, dtype="string"),
    ).fillna("No additional SHAP context available.")

    call = pd.DataFrame({
        "Customer ID": scores["customer_id_output"].fillna(""),
        "Customer": scores[CUSTOMER_COLUMN].fillna(""),
        "Customer Group": scores["customer_group_output"].fillna(""),
        "Ship-To": scores["zipcode"].astype("string").fillna(""),
        "Business Unit": scores["bulevel1"].fillna(""),
        "Branch": scores["branch_output"].fillna(""),
        "Sales Rep": scores["sales_rep_output"].fillna(""),
        "Category": scores["mgrl1"].fillna(""),
        "Risk Score": scores["category_risk_score"],
        "Risk Tier": scores["risk_tier"].fillna(""),
        "Action": scores["category_action"].fillna(""),
        "Annual Spend": scores["category_annual_spend"],
        "Rev at Risk": scores["category_revenue_at_risk"],
        "Who They Are": scores["who_they_are"],
        "Bold Signals": scores["bold_signals"],
        "Grey Signals": scores["grey_signals"],
        "Scoring Run ID": scores["scoring_run_id"].fillna(""),
    })
    call = call[CALL_LIST_COLUMNS].sort_values(
        ["Risk Score", "Annual Spend", "Customer", "Ship-To", "Category"],
        ascending=[False, False, True, True, True],
        na_position="last",
    ).reset_index(drop=True)

    cdata = scores[
        [
            CUSTOMER_COLUMN,
            "zipcode",
            "bulevel1",
            "mgrl1",
            "snapshot_dt",
            "p_14",
            "risk_tier",
            "category_action",
            "spend_365",
            "spend_365_all",
            "weighted_risk",
            "customer_risk_score",
            "customer_action",
            "revenue_at_risk",
            "scoring_run_id",
        ]
    ].copy()
    cdata["category_revenue_at_risk"] = scores[
        "category_revenue_at_risk"
    ].to_numpy()
    cdata["customer_ids"] = scores["customer_id_output"].to_numpy()
    cdata["servicingbranch"] = scores["branch_output"].to_numpy()
    cdata["salesrep"] = scores["sales_rep_output"].to_numpy()
    return call, cdata


def _style_sheet(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(color=HEADER_FONT, bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    ws.row_dimensions[1].height = 30


def _write_dataframe(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])
    _style_sheet(ws)


def _format_call_list(ws, row_count: int) -> None:
    widths = {
        "A": 18,
        "B": 32,
        "C": 18,
        "D": 12,
        "E": 14,
        "F": 14,
        "G": 20,
        "H": 26,
        "I": 12,
        "J": 12,
        "K": 16,
        "L": 16,
        "M": 16,
        "N": 60,
        "O": 70,
        "P": 70,
        "Q": 30,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2, max_row=row_count + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in ws["O"][1:]:
        cell.font = Font(bold=True)
    for cell in ws["P"][1:]:
        cell.font = Font(color=GREY_FONT)

    for col in ["L", "M"]:
        for cell in ws[col][1:]:
            cell.number_format = "$#,##0"

    if row_count:
        risk_range = f"I2:I{row_count + 1}"
        ws.conditional_formatting.add(
            risk_range,
            CellIsRule(
                operator="greaterThanOrEqual",
                formula=["40"],
                fill=PatternFill("solid", fgColor=RED_FILL),
            ),
        )
        ws.conditional_formatting.add(
            risk_range,
            CellIsRule(
                operator="between",
                formula=["20", "39"],
                fill=PatternFill("solid", fgColor=YELLOW_FILL),
            ),
        )
        ws.conditional_formatting.add(
            risk_range,
            CellIsRule(
                operator="lessThan",
                formula=["20"],
                fill=PatternFill("solid", fgColor=GREEN_FILL),
            ),
        )


def _build_main_workbook(call: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Call List"
    _write_dataframe(ws, call)
    _format_call_list(ws, len(call))

    desc = wb.create_sheet("Column Descriptions")
    desc.append(["Column", "Business Definition"])
    for name, description in COLUMN_DESCRIPTIONS:
        desc.append([name, description])
    _style_sheet(desc)
    desc.column_dimensions["A"].width = 24
    desc.column_dimensions["B"].width = 100
    for row in desc.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    lists = wb.create_sheet("Lists")
    lists.append(["Action"])
    for action in sorted(VALID_ACTIONS):
        lists.append([action])
    lists.sheet_state = "hidden"

    if len(call):
        dv = DataValidation(
            type="list",
            formula1="'Lists'!$A$2:$A$" + str(len(VALID_ACTIONS) + 1),
            allow_blank=False,
        )
        ws.add_data_validation(dv)
        dv.add(f"K2:K{len(call) + 1}")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _validation_tables(
    call: pd.DataFrame,
    reject: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    valid_actions = call["Action"].astype(str).isin(VALID_ACTIONS)
    flags = []

    def add_flags(mask: pd.Series, flag: str):
        bad = call[mask].copy()
        if bad.empty:
            return
        bad.insert(0, "Flag", flag)
        flags.append(bad)

    add_flags(~call["Risk Score"].between(0, 100), "RISK_OUT_OF_RANGE")
    add_flags(~valid_actions, "INVALID_ACTION")
    add_flags(
        pd.to_numeric(call["Rev at Risk"], errors="coerce").fillna(-1) < 0,
        "NEGATIVE_REV_AT_RISK",
    )
    add_flags(
        call["Who They Are"].astype(str).str.strip().eq(""),
        "MISSING_PROFILE",
    )
    add_flags(
        call["Customer"].astype(str).str.strip().eq(""),
        "MISSING_CUSTOMER",
    )
    add_flags(
        call["Ship-To"].astype(str).str.strip().eq(""),
        "MISSING_SHIPTO",
    )

    row_flags = (
        pd.concat(flags, ignore_index=True)
        if flags
        else pd.DataFrame(columns=["Flag", *CALL_LIST_COLUMNS])
    )

    reject_reason_ok = (
        True
        if reject.empty
        else reject["score_ineligible_reason_14"].notna().all()
    )

    reject_missing = (
        int((~reject["score_ineligible_reason_14"].notna()).sum())
        if not reject.empty
        else 0
    )
    risk_bad = int((~call["Risk Score"].between(0, 100)).sum())
    action_bad = int((~valid_actions).sum())
    rar_bad = int(
        (
            pd.to_numeric(call["Rev at Risk"], errors="coerce").fillna(-1)
            < 0
        ).sum()
    )
    profile_bad = int(
        call["Who They Are"].astype(str).str.strip().eq("").sum()
    )
    signals_bad = int(
        call["Bold Signals"].astype(str).str.strip().eq("").sum()
    )
    duplicate_bad = int(
        call.duplicated(["Customer", "Business Unit", "Ship-To", "Category"]).sum()
    )

    checks = [
        ("Layer 1", "Reject rows have a reason", reject_missing, reject_reason_ok),
        ("Layer 2", "Risk Score within [0,100]", risk_bad, risk_bad == 0),
        ("Layer 2", "Action in approved 9-tier grid", action_bad, action_bad == 0),
        ("Layer 2", "Rev at Risk non-negative", rar_bad, rar_bad == 0),
        ("Layer 2", "Who They Are populated", profile_bad, profile_bad == 0),
        ("Layer 2", "Bold Signals populated", signals_bad, signals_bad == 0),
        ("Layer 3", "No duplicate Customer x BU x ZIP x Category grain", duplicate_bad, duplicate_bad == 0),
    ]
    validation = pd.DataFrame(
        checks,
        columns=["Layer", "Check", "Failure Count", "Passed"],
    )
    validation["Status"] = np.where(
        validation["Passed"],
        "PASS",
        "FAIL",
    )
    return validation, row_flags


def _build_audit_workbook(
    path: Path,
    call: pd.DataFrame,
    attrs: pd.DataFrame,
    reject: pd.DataFrame,
    dormant: pd.DataFrame,
    validation: pd.DataFrame,
    row_flags: pd.DataFrame,
    health: pd.DataFrame,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    owner = call[
        ["Customer ID", "Customer", "Ship-To", "Sales Rep"]
    ].drop_duplicates().copy()
    owner["Account Owner"] = ""
    owner["Mapping Status"] = (
        "OFFLINE - SharePoint owner mapping not connected"
    )

    cust_id = call[
        ["Customer ID", "Customer", "Ship-To"]
    ].drop_duplicates().copy()
    cust_id["Customer ID Count"] = (
        cust_id["Customer ID"]
        .fillna("")
        .astype(str)
        .apply(
            lambda x: len(
                [p for p in x.split(",") if p.strip()]
            )
        )
    )
    cust_id["Status"] = np.select(
        [
            cust_id["Customer ID Count"].eq(0),
            cust_id["Customer ID Count"].gt(1),
        ],
        ["MISSING", "MULTIPLE"],
        default="VALID",
    )

    reject_out = reject.copy()
    if reject_out.empty:
        reject_out = pd.DataFrame(
            [{"Status": "No rejected rows for this scoring run"}]
        )
    attrs_out = (
        attrs.copy()
        if not attrs.empty
        else pd.DataFrame([{"Status": "No customer attributes"}])
    )

    sheets = [
        ("Owner Mapping Review", owner),
        ("Customer ID Validation", cust_id),
        ("Validation Layer 1-3", validation),
        ("Row Flags", row_flags),
        ("Dormant Accounts", dormant),
        ("Pipeline Health", health),
        ("Reject Exclusions", reject_out),
        ("Customer Attributes", attrs_out),
    ]

    for name, df in sheets:
        ws = wb.create_sheet(name)
        _write_dataframe(ws, df)
        for col_idx in range(1, min(ws.max_column, 20) + 1):
            max_len = 0
            for cells in ws.iter_cols(
                min_col=col_idx,
                max_col=col_idx,
                min_row=1,
                max_row=min(ws.max_row, 200),
            ):
                for cell in cells:
                    max_len = max(
                        max_len,
                        len(str(cell.value or "")),
                    )
            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = min(max(max_len + 2, 12), 60)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def generate_business_output(
    root: str | Path,
    preferred_format: str = "parquet",
) -> dict:
    root = Path(root)

    scoring = root / "data" / "scoring"
    features = root / "data" / "features"
    certified = root / "data" / "certified"
    monitoring = root / "data" / "monitoring"

    scores = read_df(
        resolve_table(
            scoring / "dex_v2_category_loss_scores_wide_zipcode"
        )
    )
    reasons = read_df(
        resolve_table(
            scoring / "dex_v2_category_reason_long_zipcode"
        )
    )
    reject = read_df(
        resolve_table(scoring / "dex_v2_weekly_reject_log")
    )
    rollup = read_df(
        resolve_table(scoring / "dex_v2_customer_zip_risk_rollup")
    )
    attrs = read_df(
        resolve_table(certified / "dex_v2_cust_attrs")
    )
    peak = read_df(
        resolve_table(
            features / "dex_v2_customer_peak_months_zipcode"
        )
    )
    top = read_df(
        resolve_table(
            features / "dex_v2_customer_top_categories_zipcode"
        )
    )

    call, cdata = _make_call_list(
        scores,
        reasons,
        attrs,
        peak,
        top,
    )

    latest_snapshot = pd.to_datetime(scores["snapshot_dt"]).max()
    date_token = latest_snapshot.strftime("%m-%d-%y")

    output_dir = root / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    main_path = output_dir / f"Weekly Call List {date_token}.xlsx"
    audit_path = (
        output_dir
        / f"Weekly Call List {date_token}_audit.xlsx"
    )

    validation, row_flags = _validation_tables(
        call,
        reject,
    )

    latest_scores = scores[
        pd.to_datetime(scores["snapshot_dt"]) == latest_snapshot
    ].copy()
    recent_spend = pd.to_numeric(
        latest_scores.get("spend_3m_g"),
        errors="coerce",
    ).fillna(0)
    dormant = latest_scores[
        recent_spend <= 0
    ][
        [
            CUSTOMER_COLUMN,
            "zipcode",
            "mgrl1",
            "p_14",
            "risk_tier",
            "spend_365",
            "spend_3m_g",
            "recency_days",
            "scoring_run_id",
        ]
    ].copy()

    drift_summary = _read_optional(
        monitoring / "dex_v2_drift_summary"
    )
    drift_status = (
        str(
            drift_summary.sort_values(
                "scoring_snapshot"
            ).iloc[-1]["overall_status"]
        )
        if not drift_summary.empty
        else "NOT_AVAILABLE"
    )

    manifest_path = (
        root / "data" / "local_pipeline_manifest.json"
    )
    manifest = (
        json.loads(
            manifest_path.read_text(encoding="utf-8")
        )
        if manifest_path.exists()
        else {}
    )
    perf_status = (
        manifest.get("monitoring", {})
        .get("performance", {})
        .get("status", "NOT_AVAILABLE")
    )

    overall_pass = bool(validation["Passed"].all())
    health = pd.DataFrame([
        {
            "Snapshot": str(latest_snapshot.date()),
            "Scoring Run ID": (
                str(
                    latest_scores[
                        "scoring_run_id"
                    ].iloc[0]
                )
                if len(latest_scores)
                else ""
            ),
            "Call List Rows": int(len(call)),
            "Scored Category Rows": int(len(latest_scores)),
            "Customer ZIP Rollup Rows": int(len(rollup)),
            "Rejected Rows": int(len(reject)),
            "Total Annual Spend": float(
                pd.to_numeric(
                    call["Annual Spend"],
                    errors="coerce",
                ).fillna(0).sum()
            ),
            "Total Rev at Risk": float(
                pd.to_numeric(
                    call["Rev at Risk"],
                    errors="coerce",
                ).fillna(0).sum()
            ),
            "Drift Status": drift_status,
            "Performance Status": perf_status,
            "Output Validation": (
                "PASS" if overall_pass else "FAIL"
            ),
            "Account Owner Mapping": (
                "OFFLINE - SharePoint mapping unavailable"
            ),
        }
    ])

    _build_main_workbook(call, main_path)
    _build_audit_workbook(
        audit_path,
        call,
        attrs,
        reject,
        dormant,
        validation,
        row_flags,
        health,
    )

    cdata_dir = root / "data" / "output"
    cdata_dir.mkdir(parents=True, exist_ok=True)
    cdata_path = write_df(
        cdata,
        cdata_dir / "dex_v2_cdata_output_snapshot",
        preferred_format,
    )

    qa = {
        "status": "PASS" if overall_pass else "FAIL",
        "snapshot_dt": str(latest_snapshot.date()),
        "rows": {
            "call_list": int(len(call)),
            "scored_categories": int(len(latest_scores)),
            "customer_zip_rollup": int(len(rollup)),
            "rejected": int(len(reject)),
            "dormant_categories": int(len(dormant)),
            "row_flags": int(len(row_flags)),
        },
        "totals": {
            "annual_spend": float(
                pd.to_numeric(
                    call["Annual Spend"],
                    errors="coerce",
                ).fillna(0).sum()
            ),
            "revenue_at_risk": float(
                pd.to_numeric(
                    call["Rev at Risk"],
                    errors="coerce",
                ).fillna(0).sum()
            ),
        },
        "validation": {
            str(r["Check"]): str(r["Status"])
            for _, r in validation.iterrows()
        },
        "paths": {
            "main_workbook": str(main_path),
            "audit_workbook": str(audit_path),
            "cdata_snapshot": str(cdata_path),
        },
    }

    qa_dir = root / "data" / "qa"
    qa_dir.mkdir(parents=True, exist_ok=True)
    (
        qa_dir / "business_output_qa.json"
    ).write_text(
        json.dumps(qa, indent=2),
        encoding="utf-8",
    )
    update_manifest(root, "business_output", qa)

    if not overall_pass:
        raise RuntimeError(
            "Business output validation failed. "
            "Review data/qa/business_output_qa.json"
        )

    return qa
